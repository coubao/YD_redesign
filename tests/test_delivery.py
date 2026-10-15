import re
import hashlib
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path

from alembic import command
from alembic.config import Config
from openpyxl import Workbook
from sqlalchemy import func, select
from werkzeug.security import generate_password_hash

from app import app
from delivery.importer import import_master, preview, rollback_import
from delivery.models import AdminEvent, ImportBatch, ParentAccess, Project, Student, Task, TaskChange, TemplateTask, User
from delivery.storage import engine, session as db_session


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.db_path = Path(self.temp.name) / 'delivery.db'
        app.config.update(TESTING=True, SECRET_KEY='test-delivery-secret-key-long-enough-123456',
                          DELIVERY_DATABASE_URL=f'sqlite:///{self.db_path}')
        config = Config(str(Path(__file__).resolve().parents[1] / 'delivery' / 'alembic.ini'))
        config.set_main_option('script_location', str(Path(__file__).resolve().parents[1] / 'delivery' / 'migrations'))
        config.set_main_option('sqlalchemy.url', f'sqlite:///{self.db_path}')
        command.upgrade(config, 'head')
        self.client = app.test_client()
        with app.app_context():
            with engine().connect() as connection:
                tables = connection.exec_driver_sql("SELECT name FROM sqlite_master WHERE type='table'").all()
                self.assertIn(('delivery_task',), tables)
            with app.test_request_context('/delivery/'):
                db = db_session()
                self.admin = User(username='admin-test', password_hash=generate_password_hash('admin-password-123'), role='admin')
                self.staff = User(username='staff-test', password_hash=generate_password_hash('staff-password-123'), role='staff')
                self.parent = User(username='parent-test', password_hash=generate_password_hash('parent-password-123'), role='parent')
                db.add_all([self.admin, self.staff, self.parent])
                db.flush()
                first = Student(name='测试甲')
                second = Student(name='测试乙')
                db.add_all([first, second])
                db.flush()
                first_project = Project(student_id=first.id, service_type='本科', manager_id=self.staff.id, current_phase='申请准备')
                second_project = Project(student_id=second.id, service_type='硕士', current_phase='审核中')
                db.add_all([first_project, second_project])
                db.flush()
                db.add(ParentAccess(user_id=self.parent.id, student_id=first.id))
                public = Task(project_id=first_project.id, name='公开任务', status='进行中', parent_visible=True,
                              public_note='家长可见说明', internal_note='绝对内部信息')
                private = Task(project_id=first_project.id, name='内部任务绝不公开', status='进行中', parent_visible=False,
                               internal_note='内部审查意见')
                db.add_all([public, private])
                db.commit()
                self.first_id, self.second_id, self.public_id = first_project.id, second_project.id, public.id

    def csrf(self, client=None):
        client = client or self.client
        html = client.get('/delivery/login').get_data(as_text=True)
        return re.search(r'name="_csrf" value="([^"]+)"', html).group(1)

    def login(self, username, password, client=None):
        client = client or self.client
        token = self.csrf(client)
        return client.post('/delivery/login', data={'_csrf': token, 'username': username, 'password': password}, follow_redirects=True)

    def token_from(self, path, client=None):
        client = client or self.client
        html = client.get(path).get_data(as_text=True)
        return re.search(r'name="_csrf" value="([^"]+)"', html).group(1)

    def test_parent_isolation_and_internal_redaction(self):
        response = self.login('parent-test', 'parent-password-123')
        self.assertEqual(response.status_code, 200)
        own = self.client.get(f'/delivery/projects/{self.first_id}')
        html = own.get_data(as_text=True)
        self.assertIn('家长可见说明', html)
        self.assertNotIn('内部任务绝不公开', html)
        self.assertNotIn('绝对内部信息', html)
        self.assertEqual(self.client.get(f'/delivery/projects/{self.second_id}').status_code, 404)
        self.assertEqual(self.client.get('/delivery/users').status_code, 403)
        self.assertEqual(self.client.get(f'/delivery/progress/{self.first_id}.txt').status_code, 403)

    def test_staff_update_reaches_parent_without_duplicate_history(self):
        self.login('staff-test', 'staff-password-123')
        project_url = f'/delivery/projects/{self.first_id}'
        token = self.token_from(project_url)
        endpoint = f'/delivery/tasks/{self.public_id}'
        self.assertEqual(self.client.post(endpoint, data={'action': 'complete'}).status_code, 400)
        response = self.client.post(endpoint, data={'_csrf': token, 'action': 'complete'})
        self.assertEqual(response.status_code, 302)
        self.client.post(endpoint, data={'_csrf': token, 'action': 'complete'})
        with app.test_request_context('/delivery/'):
            db = db_session()
            self.assertEqual(db.scalar(select(func.count()).select_from(TaskChange)), 1)
        parent_client = app.test_client()
        self.login('parent-test', 'parent-password-123', parent_client)
        html = parent_client.get(project_url).get_data(as_text=True)
        self.assertIn('已经完成', html)
        self.assertIn('家长可见说明', html)

    def test_staff_cannot_open_unassigned_project(self):
        self.login('staff-test', 'staff-password-123')
        self.assertEqual(self.client.get(f'/delivery/projects/{self.second_id}').status_code, 404)

    def test_disabled_login_and_logout(self):
        with app.test_request_context('/delivery/'):
            db = db_session()
            parent = db.scalar(select(User).where(User.username == 'parent-test'))
            parent.active = False
            db.commit()
        result = self.login('parent-test', 'parent-password-123')
        self.assertIn('账号或密码不正确', result.get_data(as_text=True))
        self.login('admin-test', 'admin-password-123')
        token = self.token_from('/delivery/')
        self.assertEqual(self.client.post('/delivery/logout', data={'_csrf': token}).status_code, 302)
        self.assertEqual(self.client.get('/delivery/').status_code, 302)

    def test_unconfigured_secret_fails_closed(self):
        app.config['SECRET_KEY'] = 'change-me-in-production'
        self.assertEqual(self.client.get('/delivery/login').status_code, 503)

    def test_production_refuses_sqlite_delivery_database(self):
        with patch.dict('os.environ', {'DELIVERY_ENV': 'production'}):
            self.assertEqual(self.client.get('/delivery/login').status_code, 503)

    def test_delivery_and_static_requests_skip_website_schema_patching(self):
        with patch('app.ensure_ranking_schema', side_effect=AssertionError('legacy database touched')):
            self.assertEqual(self.client.get('/delivery/login').status_code, 200)
            response = self.client.get('/static/delivery.css')
            self.assertEqual(response.status_code, 200)
            response.close()

    def test_admin_account_and_access_actions_are_audited(self):
        self.login('admin-test', 'admin-password-123')
        users_url = '/delivery/users'
        token = self.token_from(users_url)
        self.assertEqual(self.client.post(users_url, data={
            '_csrf': token, 'username': 'new-parent', 'password': 'some-long-password-123', 'role': 'parent',
        }).status_code, 302)
        with app.test_request_context('/delivery/'):
            db = db_session()
            new_parent = db.scalar(select(User).where(User.username == 'new-parent'))
            student = db.scalar(select(Student).where(Student.name == '测试甲'))
            parent_id, student_id = new_parent.id, student.id
        token = self.token_from(users_url)
        self.assertEqual(self.client.post('/delivery/parents/access', data={
            '_csrf': token, 'user_id': str(parent_id), 'student_id': str(student_id),
        }).status_code, 302)
        with app.test_request_context('/delivery/'):
            db = db_session()
            actions = set(db.scalars(select(AdminEvent.action)).all())
            self.assertTrue({'user_created', 'parent_access_granted'} <= actions)

    def test_pasted_text_match_requires_manual_confirmation(self):
        with app.test_request_context('/delivery/'):
            db = db_session()
            task = db.get(Task, self.public_id)
            task.source_code = 'US-01'
            db.commit()
        self.login('staff-test', 'staff-password-123')
        token = self.token_from(f'/delivery/projects/{self.first_id}')
        response = self.client.post(f'/delivery/projects/{self.first_id}/match-text', data={
            '_csrf': token, 'source_text': 'US-01 已完成',
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn('公开任务', response.get_data(as_text=True))
        with app.test_request_context('/delivery/'):
            self.assertEqual(db_session().get(Task, self.public_id).status, '进行中')

    def test_ambiguous_import_status_stays_private_until_review(self):
        with app.test_request_context('/delivery/'):
            db = db_session()
            task = db.get(Task, self.public_id)
            task.review_required = True
            task.source_status = '待确认'
            db.commit()
        self.login('parent-test', 'parent-password-123')
        html = self.client.get(f'/delivery/projects/{self.first_id}').get_data(as_text=True)
        self.assertNotIn('家长可见说明', html)
        staff_client = app.test_client()
        self.login('staff-test', 'staff-password-123', staff_client)
        token = self.token_from(f'/delivery/projects/{self.first_id}', staff_client)
        response = staff_client.post(f'/delivery/tasks/{self.public_id}', data={
            '_csrf': token, 'action': 'quick_status', 'status': '进行中',
        })
        self.assertEqual(response.status_code, 302)
        html = self.client.get(f'/delivery/projects/{self.first_id}').get_data(as_text=True)
        self.assertIn('家长可见说明', html)

    def test_reviewed_import_and_rollback(self):
        source = Path(self.temp.name) / 'synthetic_master.xlsx'
        workbook = Workbook()
        student_sheet = workbook.active
        student_sheet.title = '01_学生总表'
        student_sheet.append(['学生姓名', '产品类型', '入学季', '当前阶段', '风险等级'])
        student_sheet.append(['测试丙', '本科申请', '2028', '准备中', '绿'])
        task_sheet = workbook.create_sheet('02_申请任务总表')
        task_sheet.append(['学生姓名', '模块编号', '任务名称', '所属阶段', '实际负责人', '任务状态', '风险等级', '截止日期', '备注'])
        task_sheet.append(['测试丙', 'US-01', '测试任务', '准备中', '老师', '待确认', '黄', None, '仅内部'])
        for sheet_name in ('03_24模块模板', '04_PhD模块模板', '05_硕士模块模板', '06_硕博联申模块模板'):
            sheet = workbook.create_sheet(sheet_name)
            sheet.append(['模块编号', '任务名称', '所属阶段', '标准负责人角色'])
            sheet.append([sheet_name[:2] + '-01', '模板任务', '准备中', '顾问'])
        workbook.save(source)
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        self.assertEqual(preview([source])['sources'][0]['task_count'], 1)
        with app.test_request_context('/delivery/'):
            db = db_session()
            with self.assertRaises(ValueError):
                import_master(db, source, '0' * 64)
            self.assertEqual(db.scalar(select(func.count()).select_from(ImportBatch)), 0)
            # The import target must be empty, so delete synthetic fixtures first.
            for model in (Task, ParentAccess, Project, Student):
                for row in db.scalars(select(model)):
                    db.delete(row)
                db.flush()
            db.commit()
            batch_id, student_count, task_count, template_count = import_master(db, source, digest)
            self.assertEqual((student_count, task_count, template_count), (1, 1, 4))
            task = db.scalar(select(Task))
            self.assertEqual(task.source_status, '待确认')
            self.assertEqual(task.status, '未开始')
            self.assertTrue(task.review_required)
            self.assertFalse(task.parent_visible)
            self.assertEqual(db.scalar(select(TemplateTask)).classification, 'pending_review')
            db.add(TaskChange(task_id=task.id, actor_id=self.admin.id, changes_json='{}'))
            db.commit()
            with self.assertRaises(ValueError):
                rollback_import(db, batch_id)
            db.delete(db.scalar(select(TaskChange)))
            db.commit()
            rollback_import(db, batch_id)
            self.assertEqual(db.scalar(select(func.count()).select_from(Student)), 0)
            self.assertEqual(db.scalar(select(func.count()).select_from(TemplateTask)), 0)


if __name__ == '__main__':
    unittest.main()
