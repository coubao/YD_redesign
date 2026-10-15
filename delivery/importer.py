"""Read-only source inspection and explicit, transactional master import."""

import hashlib
import json
from datetime import date, datetime
from pathlib import Path

from sqlalchemy import func, select

from .models import ImportBatch, Project, Student, Task, TemplateTask


STUDENT_SHEET = '01_学生总表'
TASK_SHEET = '02_申请任务总表'
TEMPLATE_SHEETS = {
    '03_24模块模板': '本科',
    '04_PhD模块模板': '博士',
    '05_硕士模块模板': '硕士',
    '06_硕博联申模块模板': '硕博联申',
}
SERVICE_MAP = {'本科申请': '本科', '硕士申请': '硕士', '博士申请': '博士', '硕博联申': '硕博联申'}
STATUS_MAP = {'已完成': '已完成', '进行中': '进行中', '未开始': '未开始',
              '待触发': '未开始', '待确认': '未开始'}


def open_workbook(path):
    from openpyxl import load_workbook
    return load_workbook(path, read_only=True, data_only=True)


def sheet_records(sheet):
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else '' for value in next(rows)]
    for row_number, values in enumerate(rows, start=2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        yield row_number, dict(zip(headers, values))


def text(value):
    return str(value).strip() if value is not None else ''


def iso_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    try:
        return date.fromisoformat(text(value)[:10])
    except ValueError:
        return None


def task_index(path):
    workbook = open_workbook(path)
    indexed = {}
    duplicate = []
    for sheet_name in workbook.sheetnames:
        if '任务总表' not in sheet_name:
            continue
        for row, record in sheet_records(workbook[sheet_name]):
            code = text(record.get('模块编号'))
            student = text(record.get('学生姓名'))
            if not code or not student:
                continue
            key = (student, code)
            if key in indexed:
                duplicate.append({'sheet': sheet_name, 'row': row, 'code': code})
            indexed[key] = {'status': text(record.get('任务状态')), 'sheet': sheet_name, 'row': row}
    workbook.close()
    return indexed, duplicate


def preview(paths):
    """Return no student names or private notes in the report."""
    reports = []
    indices = []
    student_sets = []
    for raw_path in paths:
        path = Path(raw_path)
        if path.suffix.lower() == '.xlsx':
            workbook = open_workbook(path)
            sheets = {name: sum(1 for _ in sheet_records(workbook[name])) for name in workbook.sheetnames}
            student_rows = list(sheet_records(workbook[STUDENT_SHEET])) if STUDENT_SHEET in workbook.sheetnames else []
            students = [text(record.get('学生姓名')) for _, record in student_rows if text(record.get('学生姓名'))]
            seen = set()
            duplicate_student_rows = []
            for row_number, record in student_rows:
                name = text(record.get('学生姓名'))
                if name in seen:
                    duplicate_student_rows.append(row_number)
                seen.add(name)
            workbook.close()
            index, duplicate = task_index(path)
            reports.append({'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                            'sheets': sheets, 'task_count': len(index), 'duplicate_codes_within_file': duplicate,
                            'duplicate_student_rows': duplicate_student_rows})
            indices.append(index)
            student_sets.append(set(students))
        elif path.suffix.lower() == '.html':
            source = path.read_text(encoding='utf-8')
            start = source.index('const SEED=') + len('const SEED=')
            seed, _ = json.JSONDecoder().raw_decode(source[start:].lstrip())
            tasks = seed.get('tasks', [])
            index = {(text(item.get('student')), text(item.get('id'))): {'status': text(item.get('status'))}
                     for item in tasks}
            reports.append({'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                            'student_count': len(seed.get('students', [])), 'task_count': len(tasks),
                            'template_groups': len(seed.get('templates', {}))})
            indices.append(index)
            student_sets.append({text(item.get('name')) for item in seed.get('students', [])})
        else:
            raise ValueError(f'Unsupported source type: {path.suffix}')
    overlaps = []
    for left in range(len(indices)):
        for right in range(left + 1, len(indices)):
            common = indices[left].keys() & indices[right].keys()
            conflicts = [code for student, code in common
                         if indices[left][(student, code)]['status'] != indices[right][(student, code)]['status']]
            overlaps.append({'files': [reports[left]['file'], reports[right]['file']],
                             'shared_tasks': len(common), 'status_conflict_count': len(conflicts),
                             'shared_students': len(student_sets[left] & student_sets[right]),
                             'conflicting_codes': sorted(set(conflicts))})
    return {'sources': reports, 'overlaps': overlaps,
            'note': 'Only the selected master workbook may be imported; all template classifications require review.'}


def load_master(path):
    workbook = open_workbook(path)
    required = {STUDENT_SHEET, TASK_SHEET, *TEMPLATE_SHEETS}
    if not required.issubset(workbook.sheetnames):
        raise ValueError('The selected workbook is not the V1.1 master structure')
    students = list(sheet_records(workbook[STUDENT_SHEET]))
    tasks = list(sheet_records(workbook[TASK_SHEET]))
    templates = {name: list(sheet_records(workbook[name])) for name in TEMPLATE_SHEETS}
    workbook.close()
    names = [text(row.get('学生姓名')) for _, row in students]
    if not names or len(names) != len(set(names)) or not all(names):
        raise ValueError('Duplicate or blank students require manual resolution')
    keys = [(text(row.get('学生姓名')), text(row.get('模块编号'))) for _, row in tasks]
    if len(keys) != len(set(keys)) or any(not a or not b for a, b in keys):
        raise ValueError('Duplicate or blank task identifiers require manual resolution')
    if set(a for a, _ in keys) - set(names):
        raise ValueError('Task references a student missing from the master student sheet')
    if any(text(row.get('产品类型')) not in SERVICE_MAP for _, row in students):
        raise ValueError('Unknown service type requires manual resolution')
    if any(text(row.get('任务状态')) not in STATUS_MAP for _, row in tasks):
        raise ValueError('Unknown task status requires manual resolution')
    return students, tasks, templates


def import_master(db, path, expected_sha):
    path = Path(path)
    actual_sha = hashlib.sha256(path.read_bytes()).hexdigest()
    if actual_sha != expected_sha:
        raise ValueError('Source SHA-256 changed since preview')
    students, tasks, templates = load_master(path)
    if db.scalar(select(func.count()).select_from(Student)) or db.scalar(select(func.count()).select_from(TemplateTask)):
        raise ValueError('Target database already contains delivery records; no silent merge is allowed')
    if db.scalar(select(ImportBatch).where(ImportBatch.source_sha256 == actual_sha)):
        raise ValueError('This workbook was imported already')
    try:
        batch = ImportBatch(source_sha256=actual_sha, source_file=path.name,
                            student_count=len(students), task_count=len(tasks))
        db.add(batch)
        db.flush()
        projects = {}
        for row_number, record in students:
            name = text(record['学生姓名'])
            student = Student(name=name, import_batch_id=batch.id)
            db.add(student)
            db.flush()
            project = Project(student_id=student.id, import_batch_id=batch.id,
                              service_type=SERVICE_MAP[text(record['产品类型'])],
                              intake_year=text(record.get('入学季'))[:32],
                              current_phase=text(record.get('当前阶段'))[:120] or '待确认',
                              risk_level=text(record.get('风险等级'))[:24] or '待确认',
                              source_status=text(record.get('项目状态'))[:40],
                              source_file=path.name, source_sheet=STUDENT_SHEET, source_row=row_number,
                              source_payload=json.dumps(record, ensure_ascii=False, default=str))
            db.add(project)
            db.flush()
            projects[name] = project
        for row_number, record in tasks:
            project = projects[text(record['学生姓名'])]
            source_status = text(record['任务状态'])
            db.add(Task(project_id=project.id, import_batch_id=batch.id,
                        source_code=text(record['模块编号'])[:40], name=text(record['任务名称'])[:250],
                        phase=text(record['所属阶段'])[:120], owner_label=text(record.get('实际负责人'))[:120],
                        status=STATUS_MAP[source_status], source_status=source_status,
                        review_required=source_status in {'待触发', '待确认'},
                        risk_level=text(record.get('风险等级'))[:24] or '待确认',
                        due_date=iso_date(record.get('截止日期')), parent_visible=False,
                        internal_note=text(record.get('备注')),
                        source_file=path.name, source_sheet=TASK_SHEET, source_row=row_number,
                        source_payload=json.dumps(record, ensure_ascii=False, default=str)))
        for sheet_name, rows in templates.items():
            service_type = TEMPLATE_SHEETS[sheet_name]
            for row_number, record in rows:
                code = text(record.get('模块编号'))
                if not code:
                    continue
                db.add(TemplateTask(service_type=service_type, source_code=code[:40],
                                    name=text(record.get('任务名称'))[:250], phase=text(record.get('所属阶段'))[:120],
                                    role_label=text(record.get('标准负责人角色'))[:120],
                                    classification='pending_review', source_file=path.name,
                                    source_sheet=sheet_name, source_row=row_number, import_batch_id=batch.id))
        db.commit()
        return batch.id, len(students), len(tasks), sum(len(items) for items in templates.values())
    except Exception:
        db.rollback()
        raise


def rollback_import(db, batch_id):
    from .models import AdminEvent, Milestone, ParentAccess, TaskChange
    batch = db.get(ImportBatch, batch_id)
    if not batch:
        raise ValueError('Import batch not found')
    project_ids = db.scalars(select(Project.id).where(Project.import_batch_id == batch_id)).all()
    student_ids = db.scalars(select(Student.id).where(Student.import_batch_id == batch_id)).all()
    task_ids = db.scalars(select(Task.id).where(Task.import_batch_id == batch_id)).all()
    if db.scalar(select(func.count()).select_from(TaskChange).where(TaskChange.task_id.in_(task_ids))) or \
       db.scalar(select(func.count()).select_from(ParentAccess).where(ParentAccess.student_id.in_(student_ids))) or \
       db.scalar(select(func.count()).select_from(Milestone).where(Milestone.project_id.in_(project_ids))):
        raise ValueError('Import has later edits or access grants; review manually before rollback')
    if db.scalar(select(func.count()).select_from(AdminEvent).where(
        (AdminEvent.entity == 'project') & AdminEvent.entity_id.in_(project_ids) |
        (AdminEvent.entity == 'task') & AdminEvent.entity_id.in_(task_ids) |
        (AdminEvent.entity == 'student') & AdminEvent.entity_id.in_(student_ids))):
        raise ValueError('Import has later administrative changes; review manually before rollback')
    if db.scalar(select(func.count()).select_from(Task).where(Task.project_id.in_(project_ids), Task.import_batch_id.is_(None))):
        raise ValueError('Import projects contain later tasks; review manually before rollback')
    try:
        for model, field in ((Task, Task.import_batch_id), (TemplateTask, TemplateTask.import_batch_id),
                             (Project, Project.import_batch_id), (Student, Student.import_batch_id)):
            for item in db.scalars(select(model).where(field == batch_id)):
                db.delete(item)
        db.flush()
        db.delete(batch)
        db.commit()
    except Exception:
        db.rollback()
        raise
